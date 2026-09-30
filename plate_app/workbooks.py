"""Read both historical screening layouts and export the common 14-column schema."""
from datetime import date, datetime
from io import BytesIO
import re
from zipfile import ZipFile, BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FIELDS = ['ID', 'Cell_Line_Stock', 'Gene_intron', 'Cell_type', 'Conditions',
          'Primary Antibody1', 'Secondary Antibody1', 'Primary Antibody2',
          'Secondary Antibody2', 'Hoechst33342', 'Setting', 'Note']
HEADERS = ['Plate', 'Location'] + FIELDS
LOCATIONS = [f'{r}{c}' for c in range(1, 13) for r in 'ABCDEFGH']
ALIASES = {'Primary Antibody': 'Primary Antibody1', 'Secondary Antibody': 'Secondary Antibody1'}
META_FIELDS = ['experimenter', 'seeding_date', 'treatment_date', 'fixation_date',
               'fixation', 'storage', 'notes', 'elab_experiment_id']


def text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def meaningful(value):
    return text(value) not in ('', '-')


def blank_wells():
    return {loc: {key: '' for key in FIELDS} for loc in LOCATIONS}


def parse_workbook(raw):
    try:
        with ZipFile(BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError('The uncompressed Excel file exceeds 80 MB. Split the workbook into smaller files.')
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError('Select a valid .xlsx file.') from exc
    plates, seen, warnings = {}, set(), []
    metadata = {}
    try:
        for ws in wb:
            # Some valid XLSX writers omit the optional worksheet dimension.
            if ws.max_row is None or ws.max_column is None:
                ws.calculate_dimension(force=True)
            if ws.title == '_Sample_Metadata':
                rows = ws.iter_rows(values_only=True)
                heads = [text(v) for v in next(rows, ())]
                for row in rows:
                    item = dict(zip(heads, row))
                    metadata[text(item.get('Plate'))] = {
                        k: text(item[k].date()) if k.endswith('_date') and isinstance(item.get(k), datetime)
                        else text(item.get(k)) for k in META_FIELDS
                    }
                continue
            # The downloadable single-sheet template is also a valid upload.
            # Keep ignoring reference templates alongside historical plate sheets.
            if ws.title.lower() == 'template' and any(s.title.lower() not in ('template', '_sample_metadata') for s in wb):
                continue
            if ws.max_row > 20000 or ws.max_column > 256:
                raise ValueError(f'{ws.title}: Row or column limit exceeded.')
            rows = ws.iter_rows(values_only=True)
            header = [ALIASES.get(text(v), text(v)) for v in next(rows, ())]
            if not all(k in header for k in ('Plate', 'Location', 'ID', 'Cell_Line_Stock')):
                warnings.append(f'{ws.title}: Skipped because this sheet is not a plate table.')
                continue
            indices = {k: header.index(k) for k in HEADERS if k in header}
            map_start = next((i for i in range(len(header)-11)
                              if header[i:i+12] == [str(n) for n in range(1, 13)]), None)
            matrix, sheet_plates = {}, set()
            for rownum, row in enumerate(rows, 2):
                if map_start is not None and rownum <= 9:
                    matrix.update({f'{"ABCDEFGH"[rownum-2]}{c+1}': text(row[map_start+c])
                                   for c in range(12) if map_start+c < len(row)})
                item = {k: text(row[i]) if i < len(row) else '' for k, i in indices.items()}
                if not item.get('Plate') and not item.get('Location'):
                    continue
                plate, loc = item.get('Plate', ''), item.get('Location', '').upper()
                if not plate or loc not in LOCATIONS:
                    raise ValueError(f'{ws.title} row {rownum}: Invalid Plate or Location ({plate}, {loc}).')
                if (plate, loc) in seen:
                    raise ValueError(f'{ws.title} row {rownum}: Duplicate {plate} / {loc}.')
                seen.add((plate, loc))
                sheet_plates.add(plate)
                p = plates.setdefault(plate, {'plate': plate, 'wells': blank_wells(), 'metadata': {},
                                              'legacy_map': {}, 'warnings': []})
                p['wells'][loc] = {k: item.get(k, '') for k in FIELDS}
            if matrix and len(sheet_plates) == 1:
                p = plates[next(iter(sheet_plates))]
                p['legacy_map'] = matrix
                for loc, value in matrix.items():
                    listed = p['wells'][loc]['ID']
                    if (meaningful(value) or meaningful(listed)) and value != listed:
                        p['warnings'].append(f'{loc}: Table ID [{listed or "empty"}] and map label [{value or "empty"}] differ.')
        if not plates:
            raise ValueError('No importable plates found. Plate / Location / ID / Cell_Line_Stock columns are required.')
        for name, p in plates.items():
            p['metadata'] = metadata.get(name, {})
            p['used'] = sum(any(meaningful(v) for v in w.values()) for w in p['wells'].values())
            count = sum(plate == name for plate, _ in seen)
            if count != 96:
                p['warnings'].append(f'The table contains {count} wells. Missing locations have been filled with blank values.')
        return {'plates': list(plates.values()), 'warnings': warnings}
    finally:
        wb.close()


def put(ws, row, col, value):
    cell = ws.cell(row, col, value)
    # User-entered strings must never become Excel formulas.
    if isinstance(value, str):
        cell.data_type = 's'
    return cell


def style_table(ws, columns, freeze_at='C2'):
    ws.freeze_panes = freeze_at
    ws.auto_filter.ref = f'A1:{get_column_letter(columns)}{ws.max_row}'
    for c in ws[1][:columns]:
        c.fill = PatternFill('solid', fgColor='163D38')
        c.font = Font(color='FFFFFF', bold=True)
    ws.row_dimensions[1].height = 26
    for col in range(1, columns+1):
        ws.column_dimensions[get_column_letter(col)].width = 24 if col > 2 else 14


def export_metadata(plates, approvals, exported_at):
    """Export the complete current registry, with one row per well including blanks."""
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        for r, values in enumerate([headers] + rows, 1):
            for c, value in enumerate(values, 1):
                put(ws, r, c, value)
        style_table(ws, len(headers), freeze_at='C2' if name == 'All_Wells' else 'A2')
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        return ws

    def elab(p):
        eid = p['metadata'].get('elab_experiment_id', '')
        return 'https://makimono.elab.one/experiments.php?mode=view&id='+eid if eid else ''

    sheet('All_Wells', HEADERS + META_FIELDS + ['eLabFTW_URL', 'plate_id', 'version', 'updated_at'], [
        [p['plate'], loc] + [p['wells'][loc].get(k, '') for k in FIELDS] +
        [p['metadata'].get(k, '') for k in META_FIELDS] + [elab(p), p['id'], p['version'], p['updated_at']]
        for p in plates for loc in LOCATIONS])
    sheet('Plates', ['Plate'] + META_FIELDS + ['eLabFTW_URL', 'plate_id', 'version', 'updated_at', 'source_id', 'source_plate'], [
        [p['plate']] + [p['metadata'].get(k, '') for k in META_FIELDS] +
        [elab(p), p['id'], p['version'], p['updated_at'], p.get('source_id', ''), p.get('source_plate', '')]
        for p in plates])
    keys = ['field', 'value', 'reason', 'approved_by', 'approved_at', 'enabled', 'id', 'version']
    sheet('Approved_Labels', keys, [[a.get(k, '') for k in keys] for a in approvals])
    sheet('Export_Info', ['Field', 'Details'], [
        ['Exported at (UTC)', exported_at], ['Scope', 'Latest saved versions of all active plates, regardless of search filters.'],
        ['All_Wells', 'All 96 wells per plate, including blanks. Plate metadata is included on every row.'],
        ['Plates', 'Plate metadata, IDs, saved versions, update times and source references.'],
        ['Approved_Labels', 'Shared approvals, including enabled and disabled entries. enabled: 1 = enabled, 0 = disabled.'],
        ['Excluded', 'Unsaved changes, Trash, previous revisions and uploaded original files.'],
        ['Reimport', 'Well values in All_Wells can be imported. Shared approvals, internal IDs and versions are not restored. Back up the database for a complete backup.'],
    ])
    # Preserve metadata round-tripping through the existing importer.
    sheet('_Sample_Metadata', ['Plate'] + META_FIELDS,
          [[p['plate']] + [p['metadata'].get(k, '') for k in META_FIELDS] for p in plates])
    wb['_Sample_Metadata'].sheet_state = 'hidden'
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def export_workbook(plates, maps=False):
    wb = Workbook()
    wb.remove(wb.active)
    table = None
    if not maps:
        table = wb.create_sheet('Sheet1')
        table.append(HEADERS)
    for number, p in enumerate(plates, 1):
        ws = wb.create_sheet(f'P{number}') if maps else table
        if maps:
            ws.append(HEADERS)
        for loc in LOCATIONS:
            row = [p['plate'], loc] + [p['wells'][loc].get(k, '') for k in FIELDS]
            r = ws.max_row + 1
            for c, value in enumerate(row, 1):
                put(ws, r, c, value)
        if maps:
            for c in range(1, 13):
                ws.cell(1, 16+c, c)
                ws.column_dimensions[get_column_letter(16+c)].width = 17
            for r, letter in enumerate('ABCDEFGH', 2):
                ws.cell(r, 16, letter)
                ws.row_dimensions[r].height = 36
                for c in range(1, 13):
                    cell = put(ws, r, 16+c, p['wells'][f'{letter}{c}']['ID'])
                    cell.alignment = Alignment(wrap_text=True, vertical='center')
                    cell.fill = PatternFill('solid', fgColor='E6F2EC')
            style_table(ws, len(HEADERS))
    if table is not None:
        style_table(table, len(HEADERS))
    meta = wb.create_sheet('_Sample_Metadata')
    meta.append(['Plate'] + META_FIELDS)
    for r, p in enumerate(plates, 2):
        for c, v in enumerate([p['plate']] + [p['metadata'].get(k, '') for k in META_FIELDS], 1):
            put(meta, r, c, v)
    style_table(meta, len(META_FIELDS)+1)
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
